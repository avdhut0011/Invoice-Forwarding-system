import io
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from django.shortcuts import render
from django.http import HttpResponse
from django.contrib.auth.decorators import login_required
from django.db.models import Count, Sum, Q
from apps.receipts.models import GoodsReceipt, ReceiptStatusChoices
from apps.invoices.models import Invoice, InvoiceAllocation

@login_required
def reports_summary_view(request):
    """
    Renders an interactive summary dashboard for reporting with date filtering.
    """
    start_date = request.GET.get('start_date')
    end_date = request.GET.get('end_date')

    receipts_qs = GoodsReceipt.objects.all()
    invoices_qs = Invoice.objects.all()

    if start_date:
        receipts_qs = receipts_qs.filter(material_date__gte=start_date)
        invoices_qs = invoices_qs.filter(invoice_date__gte=start_date)
    if end_date:
        receipts_qs = receipts_qs.filter(material_date__lte=end_date)
        invoices_qs = invoices_qs.filter(invoice_date__lte=end_date)

    # Store Person Summary
    store_summary = (
        receipts_qs.values('store_person_raw')
        .annotate(
            total_receipts=Count('id'),
            bill_linked=Count('id', filter=Q(status=ReceiptStatusChoices.BILL_LINKED)),
            custody_ack=Count('id', filter=Q(status=ReceiptStatusChoices.CUSTODY_ACKNOWLEDGED)),
            pending_handover=Count('id', filter=Q(status=ReceiptStatusChoices.PENDING_HANDOVER)),
        )
        .order_by('-total_receipts')
    )

    # Bill Linkage Summary
    bills = invoices_qs.prefetch_related('allocations__goods_receipt').order_by('-created_at')[:50]

    context = {
        'start_date': start_date or '',
        'end_date': end_date or '',
        'store_summary': store_summary,
        'bills': bills,
        'total_receipts_count': receipts_qs.count(),
        'total_bills_count': invoices_qs.count(),
        'total_invoice_amount': invoices_qs.aggregate(s=Sum('gross_amount'))['s'] or 0,
    }
    return render(request, 'reports/summary.html', context)

@login_required
def export_excel_report(request):
    """
    Generates a 3-tab professional Excel (.xlsx) report matching the FMS specification:
      Tab 1: Store Person-wise Summary
      Tab 2: Bill Linkage Summary
      Tab 3: Full Raw Data (A-AM Columns)
    """
    start_date = request.GET.get('start_date')
    end_date = request.GET.get('end_date')

    receipts_qs = GoodsReceipt.objects.select_related('unit', 'supplier', 'assigned_purchase_person', 'submitted_by').all()
    invoices_qs = Invoice.objects.select_related('unit', 'supplier', 'submitted_by').prefetch_related('allocations__goods_receipt').all()

    if start_date:
        receipts_qs = receipts_qs.filter(material_date__gte=start_date)
        invoices_qs = invoices_qs.filter(invoice_date__gte=start_date)
    if end_date:
        receipts_qs = receipts_qs.filter(material_date__lte=end_date)
        invoices_qs = invoices_qs.filter(invoice_date__lte=end_date)

    wb = Workbook()
    
    # Styles
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    bold_font = Font(name="Calibri", size=11, bold=True)
    regular_font = Font(name="Calibri", size=10)
    border_side = Side(border_style="thin", color="CBD5E1")
    cell_border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")

    # ── TAB 1: Store Person-wise Summary ──────────────────────────────────────
    ws1 = wb.active
    ws1.title = "Store Person Summary"
    ws1.views.sheetView[0].showGridLines = True

    ws1.append(["FMS — Store Department Performance Summary"])
    ws1.cell(row=1, column=1).font = Font(name="Calibri", size=14, bold=True, color="1E293B")
    ws1.append([f"Generated: {datetime.now().strftime('%d/%m/%Y %H:%M')} | Date Filter: {start_date or 'All'} to {end_date or 'All'}"])
    ws1.append([])

    headers1 = ["Store Representative", "Total Receipts", "Bill Linked", "Custody Confirmed", "Pending Handover"]
    ws1.append(headers1)
    hrow1 = 4
    for col_idx in range(1, len(headers1) + 1):
        cell = ws1.cell(row=hrow1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = align_center

    store_summary = (
        receipts_qs.values('store_person_raw')
        .annotate(
            total_receipts=Count('id'),
            bill_linked=Count('id', filter=Q(status=ReceiptStatusChoices.BILL_LINKED)),
            custody_ack=Count('id', filter=Q(status=ReceiptStatusChoices.CUSTODY_ACKNOWLEDGED)),
            pending_handover=Count('id', filter=Q(status=ReceiptStatusChoices.PENDING_HANDOVER)),
        )
        .order_by('-total_receipts')
    )

    for item in store_summary:
        ws1.append([
            item['store_person_raw'] or 'Unassigned',
            item['total_receipts'],
            item['bill_linked'],
            item['custody_ack'],
            item['pending_handover']
        ])
        for col_idx in range(1, 6):
            c = ws1.cell(row=ws1.max_row, column=col_idx)
            c.font = regular_font
            c.border = cell_border
            if col_idx > 1:
                c.alignment = align_center

    # ── TAB 2: Bill Linkage Summary ───────────────────────────────────────────
    ws2 = wb.create_sheet(title="Bill Linkage Summary")
    ws2.views.sheetView[0].showGridLines = True

    ws2.append(["FMS — Bill & Invoice Linkage Summary"])
    ws2.cell(row=1, column=1).font = Font(name="Calibri", size=14, bold=True, color="1E293B")
    ws2.append([f"Generated: {datetime.now().strftime('%d/%m/%Y %H:%M')}"])
    ws2.append([])

    headers2 = ["Bill Unique ID", "Bill Number", "Bill Date", "Amount (₹)", "Supplier", "Unit", "Linked GRNs", "Due Date", "Status", "Submitted By"]
    ws2.append(headers2)
    hrow2 = 4
    for col_idx in range(1, len(headers2) + 1):
        cell = ws2.cell(row=hrow2, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = align_center

    for b in invoices_qs:
        linked_grns = [a.goods_receipt.grn_number for a in b.allocations.all() if a.goods_receipt]
        grn_str = ", ".join(linked_grns) if linked_grns else (b.na_unique_id or "N/A")
        ws2.append([
            b.bill_unique_id,
            b.invoice_number,
            b.invoice_date.strftime('%d/%m/%Y') if b.invoice_date else '',
            float(b.gross_amount),
            b.supplier_name_raw,
            b.unit.name,
            grn_str,
            b.due_date.strftime('%d/%m/%Y') if b.due_date else '',
            b.get_status_display(),
            b.purchase_person_raw or b.submitted_by.username
        ])
        for col_idx in range(1, 11):
            c = ws2.cell(row=ws2.max_row, column=col_idx)
            c.font = regular_font
            c.border = cell_border
            if col_idx == 4:
                c.alignment = align_right
                c.number_format = '₹#,##0.00'

    # ── TAB 3: Full Raw Data (Columns A–AM) ────────────────────────────────────
    ws3 = wb.create_sheet(title="Full Raw Data (A-AM)")
    ws3.views.sheetView[0].showGridLines = True

    raw_headers = [
        "Timestamp (A)", "Email (B)", "Material Date (C)", "Unit (D)", "Supplier (E)",
        "Description (F)", "Doc Type (G)", "GRN Number (H)", "GRN Date (I)", "PO/WO Number (J)",
        "DC Number (K)", "Payment Category (L)", "Invoice Received (M)", "File Link (N)", "Store Remark (O)",
        "Purchase Person (P)", "Store Person (Q)", "Unique ID (R)", "GRN Scan Copy Link (S)",
        "Purchase Status (AA)", "Bill Unique ID (AB)", "Bill Number (AC)", "Bill Date (AD)",
        "Bill Amount (₹) (AE)", "Bill File URL (AF)", "Bill Submitted By (AG)", "Bill Submission TS (AH)",
        "Supplier Name (PU) (AI)", "Supplier Code (AJ)", "GRN Received From Stores On (AK)",
        "Due Date of Payment (AL)", "Purchase Remark (AM)"
    ]
    ws3.append(raw_headers)
    for col_idx in range(1, len(raw_headers) + 1):
        cell = ws3.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = align_center

    for r in receipts_qs:
        # Find linked bill if any
        linked_alloc = r.allocations.first()
        b = linked_alloc.invoice if linked_alloc else None

        ws3.append([
            r.created_at.strftime('%d/%m/%Y %H:%M:%S'),
            r.submitter_email or '',
            r.material_date.strftime('%d/%m/%Y') if r.material_date else '',
            r.unit.name,
            r.supplier_name_raw,
            r.description,
            r.get_doc_type_display(),
            r.grn_number,
            r.grn_date.strftime('%d/%m/%Y') if r.grn_date else '',
            r.po_wo_number or '',
            r.dc_number or '',
            r.payment_category,
            "Yes" if r.is_invoice_received else "No",
            r.invoice_scan.url if r.invoice_scan else '',
            r.store_remark or '',
            r.assigned_purchase_person_raw or '',
            r.store_person_raw or '',
            r.st_number,
            r.grn_scan.url if r.grn_scan else '',
            r.get_status_display(),
            b.bill_unique_id if b else '',
            b.invoice_number if b else '',
            b.invoice_date.strftime('%d/%m/%Y') if (b and b.invoice_date) else '',
            float(b.gross_amount) if b else '',
            b.bill_scan.url if (b and b.bill_scan) else '',
            b.purchase_person_raw if b else '',
            b.created_at.strftime('%d/%m/%Y %H:%M:%S') if b else '',
            b.supplier_name_raw if b else '',
            b.supplier_code if b else '',
            b.grn_received_from_stores_on.strftime('%d/%m/%Y') if (b and b.grn_received_from_stores_on) else '',
            b.due_date.strftime('%d/%m/%Y') if (b and b.due_date) else '',
            b.purchase_remark if b else '',
        ])
        for col_idx in range(1, len(raw_headers) + 1):
            c = ws3.cell(row=ws3.max_row, column=col_idx)
            c.font = regular_font
            c.border = cell_border

    # Auto-adjust column widths
    for sheet in [ws1, ws2, ws3]:
        for col in sheet.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = col[0].column_letter
            sheet.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 40)

    # Response buffer
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    filename = f"FMS_Report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    response = HttpResponse(
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response
